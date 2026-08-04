import openpyxl

FILE = "/Users/rahulmedhe/Downloads/INVENTORY SETUP/TALWAR/TALWAR PRICING.xlsx"

data = [
    ("Royal Sword",       "Metal", "Large", "Brown"),
    ("Dragon Sword",      "Metal", "Large", "White"),
    ("Mughal Sword",      "Metal", "Large", "White"),
    ("Velvet Sword",      "Metal", "Large", "Black"),
    ("Bridal Sword",      "Metal", "Large", "Red"),
    ("Rajput Sword",      "Metal", "Large", "Red"),
    ("Teal Turban Pin",   "Metal", "Free",  "Teal"),
    ("Chrome Sword",      "Metal", "Large", "Black"),
    ("Pearl Sword",       "Metal", "Large", "White"),
    ("Filigree Sword",    "Metal", "Large", "Black"),
    ("Golden Sword",      "Metal", "Large", "White"),
    ("Teak Sword",        "Metal", "Large", "Brown"),
    ("Silver Sword",      "Metal", "Large", "Silver"),
    ("Zari Sword",        "Metal", "Large", "Red"),
    ("Steel Sword",       "Metal", "Large", "Black"),
    ("Ruby Turban Pin",   "Metal", "Free",  "Ruby"),
    ("Heritage Sword",    "Metal", "Large", "Brown"),
    ("Ruby Turban Pin II","Metal", "Free",  "Ruby"),
    ("Brass Sword",       "Metal", "Large", "Brown"),
    ("Classic Sword",     "Metal", "Large", "Brown"),
    ("Lal Sword",         "Metal", "Large", "Red"),
    ("Ivory Sword",       "Metal", "Large", "White"),
    ("Dragon Sword II",   "Metal", "Large", "Silver"),
    ("Shadow Sword",      "Metal", "Large", "Black"),
    ("Guard Sword",       "Metal", "Large", "Black"),
    ("Sandal Sword",      "Metal", "Large", "Brown"),
    ("Maroon Sword",      "Metal", "Large", "Maroon"),
]

wb = openpyxl.load_workbook(FILE)
ws = wb.active

max_col = ws.max_column

ws.cell(row=1, column=max_col+1).value = "Product Name"
ws.cell(row=1, column=max_col+2).value = "Material"
ws.cell(row=1, column=max_col+3).value = "Size"
ws.cell(row=1, column=max_col+4).value = "Colour"

for i, (name, material, size, colour) in enumerate(data):
    row = i + 2
    ws.cell(row=row, column=max_col+1).value = name
    ws.cell(row=row, column=max_col+2).value = material
    ws.cell(row=row, column=max_col+3).value = size
    ws.cell(row=row, column=max_col+4).value = colour

wb.save(FILE)
print(f"Done! Updated {len(data)} rows in TALWAR PRICING.xlsx")
